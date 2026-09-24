import os
import openpyxl
import datetime
from app.extensions import db
from app.models.student import Student
from app.models.classroom import ClassSection
from app.models.subject import Subject
from app.models.assessment import Assessment
from app.models.mark import StudentMark

class MarkService:
    @staticmethod
    def get_marks_sheet(class_section_id: int, assessment_id: int, subject_id: int):
        classroom = db.session.get(ClassSection, class_section_id)
        if not classroom:
            raise ValueError("Classroom not found.")

        assessment = db.session.get(Assessment, assessment_id)
        if not assessment:
            raise ValueError("Assessment not found.")

        subject = db.session.get(Subject, subject_id)
        if not subject:
            raise ValueError("Subject not found.")

        students = Student.query.filter_by(
            class_section_id=class_section_id,
            status="ACTIVE"
        ).order_by(Student.roll_no.asc()).all()

        existing_marks = StudentMark.query.filter_by(
            class_section_id=class_section_id,
            assessment_id=assessment_id,
            subject_id=subject_id
        ).all()

        mark_map = {m.student_id: m for m in existing_marks}

        rows = []
        valid_marks = []
        absent_count = 0
        is_sheet_locked = False

        for st in students:
            m = mark_map.get(st.id)
            if m and m.is_locked:
                is_sheet_locked = True

            obtained = m.marks_obtained if m else None
            is_ab = m.is_absent if m else False

            if is_ab:
                absent_count += 1
            elif obtained is not None:
                valid_marks.append(obtained)

            rows.append({
                "student_id": st.id,
                "roll_no": st.roll_no,
                "name": st.name,
                "admission_no": st.admission_no,
                "mark_id": m.id if m else None,
                "marks_obtained": obtained,
                "max_marks": assessment.max_marks,
                "percentage": m.percentage if m else None,
                "is_absent": is_ab,
                "is_exempt": m.is_exempt if m else False,
                "is_locked": m.is_locked if m else False,
                "remarks": m.remarks if m else None
            })

        avg_marks = round(sum(valid_marks) / len(valid_marks), 2) if valid_marks else 0.0
        high_marks = max(valid_marks) if valid_marks else 0.0
        low_marks = min(valid_marks) if valid_marks else 0.0

        return {
            "classroom": classroom.to_dict(),
            "assessment": assessment.to_dict(),
            "subject": subject.to_dict(),
            "is_locked": is_sheet_locked,
            "stats": {
                "total_students": len(students),
                "evaluated_count": len(valid_marks),
                "absent_count": absent_count,
                "average_marks": avg_marks,
                "highest_marks": high_marks,
                "lowest_marks": low_marks
            },
            "rows": rows
        }

    @staticmethod
    def save_single_mark(data: dict, evaluator_id: int = None):
        student_id = data["student_id"]
        assessment_id = data["assessment_id"]
        subject_id = data["subject_id"]

        student = db.session.get(Student, student_id)
        assessment = db.session.get(Assessment, assessment_id)
        subject = db.session.get(Subject, subject_id)

        if not student or not assessment or not subject:
            raise ValueError("Student, Assessment, or Subject not found.")

        # Check lock status on existing
        mark_rec = StudentMark.query.filter_by(
            student_id=student_id,
            assessment_id=assessment_id,
            subject_id=subject_id
        ).first()

        if mark_rec and mark_rec.is_locked:
            raise ValueError("This marks record has been locked and cannot be modified.")

        is_absent = data.get("is_absent", False)
        is_exempt = data.get("is_exempt", False)

        if is_absent:
            marks_obtained = None
        else:
            raw_marks = data.get("marks_obtained")
            if raw_marks is None:
                raise ValueError("Marks obtained is required for present student.")
            marks_obtained = float(raw_marks)
            if marks_obtained > assessment.max_marks:
                raise ValueError(f"Marks obtained ({marks_obtained}) exceeds assessment maximum ({assessment.max_marks}).")
            if marks_obtained < 0:
                raise ValueError("Marks obtained cannot be negative.")

        if mark_rec:
            mark_rec.marks_obtained = marks_obtained
            mark_rec.max_marks = assessment.max_marks
            mark_rec.is_absent = is_absent
            mark_rec.is_exempt = is_exempt
            mark_rec.evaluator_id = evaluator_id or mark_rec.evaluator_id
            mark_rec.remarks = data.get("remarks", mark_rec.remarks)
        else:
            mark_rec = StudentMark(
                school_id=student.school_id,
                session_id=student.classroom.session_id if student.classroom else assessment.session_id,
                class_section_id=student.class_section_id,
                student_id=student_id,
                assessment_id=assessment_id,
                subject_id=subject_id,
                marks_obtained=marks_obtained,
                max_marks=assessment.max_marks,
                is_absent=is_absent,
                is_exempt=is_exempt,
                evaluator_id=evaluator_id,
                remarks=data.get("remarks")
            )
            db.session.add(mark_rec)

        db.session.commit()
        return mark_rec

    @staticmethod
    def bulk_save_marks(class_section_id: int, assessment_id: int, subject_id: int, entries: list, evaluator_id: int = None):
        assessment = db.session.get(Assessment, assessment_id)
        if not assessment:
            raise ValueError("Assessment not found.")

        saved_count = 0
        for entry in entries:
            student_id = entry.get("student_id")
            if not student_id:
                continue

            MarkService.save_single_mark({
                "student_id": student_id,
                "assessment_id": assessment_id,
                "subject_id": subject_id,
                "marks_obtained": entry.get("marks_obtained"),
                "is_absent": entry.get("is_absent", False),
                "is_exempt": entry.get("is_exempt", False),
                "remarks": entry.get("remarks")
            }, evaluator_id=evaluator_id)
            saved_count += 1

        return saved_count

    @staticmethod
    def lock_marks(class_section_id: int, assessment_id: int, subject_id: int, lock_state: bool = True):
        records = StudentMark.query.filter_by(
            class_section_id=class_section_id,
            assessment_id=assessment_id,
            subject_id=subject_id
        ).all()

        for r in records:
            r.is_locked = lock_state

        db.session.commit()
        return len(records)

    @staticmethod
    def seed_marks_from_workbook(class_section_id: int):
        """
        Seeds PA-1 marks from PASHEET of FLN III-A 2026-27 (1).xlsx for Class III-A.
        """
        template_path = r"A:\fln xml system\templates_storage\FLN_TEMPLATE.xlsx"
        if not os.path.exists(template_path):
            template_path = r"A:\fln xml system\FLN III-A 2026-27 (1).xlsx"

        if not os.path.exists(template_path):
            raise FileNotFoundError("Workbook template not found.")

        classroom = db.session.get(ClassSection, class_section_id)
        if not classroom:
            raise ValueError("Classroom not found.")

        pa1 = Assessment.query.filter_by(school_id=classroom.school_id, session_id=classroom.session_id, code="PA1").first()
        if not pa1:
            raise ValueError("PA1 assessment not found.")

        # Subjects
        hin = Subject.query.filter_by(school_id=classroom.school_id, code="HIN").first()
        eng = Subject.query.filter_by(school_id=classroom.school_id, code="ENG").first()
        math = Subject.query.filter_by(school_id=classroom.school_id, code="MATH").first()
        evs = Subject.query.filter_by(school_id=classroom.school_id, code="EVS").first()

        wb = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
        if "PASHEET" not in wb.sheetnames:
            wb.close()
            raise ValueError("PASHEET sheet not found.")

        ws = wb["PASHEET"]
        saved = 0

        # Row 3 format:
        # Col 0: S.No (roll_no)
        # Col 5: HINDI (PA-1)
        # Col 6: ENG (PA-1)
        # Col 7: MATHS (PA-1)
        # Col 8: EVS (PA-1)
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None or not str(row[0]).strip().isdigit():
                continue

            roll_no = int(row[0])
            st = Student.query.filter_by(class_section_id=class_section_id, roll_no=roll_no).first()
            if not st:
                continue

            # Save Hindi
            if hin and len(row) > 5 and row[5] is not None:
                val = row[5]
                is_ab = str(val).strip().lower() in ["ab", "absent", "a"]
                try:
                    m = float(val) if not is_ab else None
                except ValueError:
                    m = None
                    is_ab = True
                MarkService.save_single_mark({
                    "student_id": st.id,
                    "assessment_id": pa1.id,
                    "subject_id": hin.id,
                    "marks_obtained": m,
                    "is_absent": is_ab
                })
                saved += 1

            # Save English
            if eng and len(row) > 6 and row[6] is not None:
                val = row[6]
                is_ab = str(val).strip().lower() in ["ab", "absent", "a"]
                try:
                    m = float(val) if not is_ab else None
                except ValueError:
                    m = None
                    is_ab = True
                MarkService.save_single_mark({
                    "student_id": st.id,
                    "assessment_id": pa1.id,
                    "subject_id": eng.id,
                    "marks_obtained": m,
                    "is_absent": is_ab
                })
                saved += 1

            # Save Maths
            if math and len(row) > 7 and row[7] is not None:
                val = row[7]
                is_ab = str(val).strip().lower() in ["ab", "absent", "a"]
                try:
                    m = float(val) if not is_ab else None
                except ValueError:
                    m = None
                    is_ab = True
                MarkService.save_single_mark({
                    "student_id": st.id,
                    "assessment_id": pa1.id,
                    "subject_id": math.id,
                    "marks_obtained": m,
                    "is_absent": is_ab
                })
                saved += 1

            # Save EVS
            if evs and len(row) > 8 and row[8] is not None:
                val = row[8]
                is_ab = str(val).strip().lower() in ["ab", "absent", "a"]
                try:
                    m = float(val) if not is_ab else None
                except ValueError:
                    m = None
                    is_ab = True
                MarkService.save_single_mark({
                    "student_id": st.id,
                    "assessment_id": pa1.id,
                    "subject_id": evs.id,
                    "marks_obtained": m,
                    "is_absent": is_ab
                })
                saved += 1

        wb.close()
        return saved
